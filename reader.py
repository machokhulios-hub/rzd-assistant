"""Документы для чтения на сайте: PDF, DOCX и TXT -> HTML единого вида.

Оформление - по правилам для нормативных документов (ГОСТ Р 7.0.97): шрифт с засечками, текст по ширине,
абзацный отступ, заголовки по центру полужирным. Из PDF убираются переносы слов, колонтитулы, номера
страниц и случайные отступы - абзацы собираются заново по положению строк на странице.
Якоря «page=N» позволяют поиску открыть документ на нужной странице."""
import html
import re
from pathlib import Path

CSS = """:root{color-scheme:light}
html,body{margin:0;background:#fff;color:#111}
body{font:18px/1.5 "Times New Roman","Liberation Serif","PT Serif",Georgia,serif;-webkit-text-size-adjust:100%}
.bar{position:sticky;top:0;z-index:1;display:flex;align-items:center;gap:12px;padding:8px 16px;background:rgba(255,255,255,.94);
  border-bottom:1px solid #e3e3e3;font:600 13px/1.3 -apple-system,system-ui,"Segoe UI",Roboto,sans-serif;color:#666;-webkit-backdrop-filter:blur(8px);backdrop-filter:blur(8px)}
.bar span{flex:1;min-width:0;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
.bar a{color:#1a5fb4;text-decoration:none;white-space:nowrap}
main{max-width:720px;margin:0 auto;padding:18px 16px 64px;overflow-wrap:break-word}
p{margin:0 0 .45em;text-align:left;text-indent:1.5em;hyphens:auto;-webkit-hyphens:auto}
@media (min-width:600px){p{text-align:justify}}
p.c{text-align:center;text-indent:0}
p.r{text-align:left;text-indent:0;margin-left:45%}
p.ed{font-size:.85em;color:#555;font-style:italic;text-indent:0;margin:.2em 0 .6em}
h2,h3{text-align:center;font-weight:700;line-height:1.3;margin:1.3em 0 .6em;hyphens:none}
h2{font-size:1.1em}h3{font-size:1em}
p.cap{text-indent:0;margin-top:.8em}
h2+h2,h2+h3,h3+h3{margin-top:-.3em}
.box{border-left:3px solid #c9c9c9;background:#f6f6f6;padding:8px 12px;margin:.6em 0;font-size:.9em}
.box p{text-indent:0;text-align:left}
.tbl{overflow-x:auto;margin:.7em 0 1em;-webkit-overflow-scrolling:touch}
table{border-collapse:collapse;font-size:.82em;line-height:1.35;min-width:60%}
td,th{border:1px solid #9a9a9a;padding:4px 6px;vertical-align:top;text-align:left;min-width:4.5em}
th{background:#f2f2f2;font-weight:700}
figure{margin:.8em 0;text-align:center}figure img{max-width:100%;height:auto}
a[id]{display:block;position:relative;top:-50px;visibility:hidden}"""


def page_html(title: str, body: str, original: str = "") -> str:
    link = f'<a href="{html.escape(original)}">Оригинал</a>' if original else ""
    return (f'<!doctype html><html lang="ru"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">'
            f'<title>{html.escape(title)}</title><style>{CSS}</style></head><body>'
            f'<div class="bar"><span>{html.escape(title)}</span>{link}</div><main>\n{body}\n</main></body></html>')


# ---- общее: что считать заголовком, как склеивать строки
HEAD_RE = re.compile(r"^(?:[IVXLC]+\.?\s|Глава\s|ГЛАВА\s|Раздел\s|РАЗДЕЛ\s|Приложение\s|ПРИЛОЖЕНИЕ\s|\d+(?:\.\d+)*\.?\s+[А-ЯЁ][А-ЯЁ\s,.-]{6,}$)")
LIST_RE = re.compile(r"^(?:[-–—•]\s|[а-яa-z]\)\s|\d+(?:\.\d+)*[.)]?\s)")


def esc(t: str) -> str:
    return html.escape(t, quote=False)


def heading_level(t: str) -> str:
    return "h2" if t.isupper() or HEAD_RE.match(t) else "h3"


# ---- PDF
def _is_noise(t: str) -> bool:
    return not t.strip() or re.fullmatch(r"[\s\d\-–—./стрСтр]+", t.strip()) is not None


def _bold(s) -> bool:
    return bool(s["flags"] & 16) or any(w in s["font"] for w in ("Bold", "Black", "Heavy", "Semibold"))


def _segs_html(segs) -> str:
    """Куски текста с выделением -> HTML; соседние куски одного вида сливаются."""
    out, cur, buf = [], None, ""
    for t, b in segs + [("", None)]:
        if b != cur and buf:
            out.append(f"<b>{esc(buf)}</b>" if cur and buf.strip() else esc(buf))
            buf = ""
        cur, buf = b, buf + t
    return re.sub(r"[ \t\n\r\f\v]+", " ", "".join(out)).strip().replace(" \u2028", "\u2028").replace("\u2028 ", "\u2028").replace("\u2028", "<br>")


def _add(segs, line):
    """Дописывает строку к абзацу: перенос слова убирается, между строками - пробел."""
    if not segs:
        segs.extend(line)
        return
    if segs[-1][0] == "\u2028":
        segs.extend(line)
        return
    a, b = segs[-1][0].rstrip(), line[0][0].lstrip()
    if re.search(r"[а-яёa-z]-$", a) and re.match(r"[а-яё]", b):
        segs[-1] = (a[:-1], segs[-1][1])
    else:
        segs[-1] = (a + " ", segs[-1][1])
    segs.extend([(b, line[0][1])] + line[1:])


def _cell(c) -> str:
    c = re.sub(r"([а-яёa-z])-\n([а-яё])", r"\1\2", c or "")
    return re.sub(r"\s+", " ", c).strip()


def pdf_has_text(path: Path) -> bool:
    import pymupdf
    return any(p.get_text().strip() for p in pymupdf.open(path))


def pdf_html(path: Path, img_dir: Path, img_url: str) -> str:
    import pymupdf
    doc = pymupdf.open(path)
    pages = list(doc)
    # колонтитулы: строки у краёв листа, которые повторяются на многих страницах (цифры не в счёт)
    edge = {}
    for p in pages:
        H = p.rect.height
        for b in p.get_text("dict")["blocks"]:
            for l in b.get("lines", []):
                if l["bbox"][3] < H * .08 or l["bbox"][1] > H * .9:
                    k = re.sub(r"\d+", "#", "".join(s["text"] for s in l["spans"]).strip())
                    edge[k] = edge.get(k, 0) + 1
    repeat = {k for k, n in edge.items() if k and n >= max(3, len(pages) * .3)}

    out, para = [], None
    img_dir.mkdir(parents=True, exist_ok=True)
    nimg = 0

    def flush():
        nonlocal para
        if para:
            h = _segs_html(para["segs"])
            plain = html.unescape(re.sub(r"<[^>]+>", " ", h)).strip()
            k = para["kind"]
            if h:
                if k == "h":
                    lv = heading_level(plain)
                    out.append(f"<{lv}>{esc(plain)}</{lv}>")
                elif k in ("c", "r", "ed"):
                    out.append(f'<p class="{k}">{h}</p>')
                else:     # «Таблица N - название» - без абзацного отступа, как требует ГОСТ
                    out.append(f'<p class="cap">{h}</p>' if re.match(r"(Таблица|Рисунок)\s+[\dIVXА-Я]", plain) else f"<p>{h}</p>")
        para = None

    for pn, p in enumerate(pages, 1):
        W, H = p.rect.width, p.rect.height
        # таблицы: настоящие рисуем таблицей, рамку из одной ячейки - врезкой
        tabs = []
        try:
            for t in p.find_tables().tables:
                rows = [[_cell(c) for c in r] for r in t.extract()]
                rows = [r for r in rows if any(r)]
                keep = [j for j in range(max(map(len, rows), default=0)) if any(j < len(r) and r[j] for r in rows)]
                rows = [[r[j] if j < len(r) else "" for j in keep] for r in rows]    # пустые столбцы - это рамка, не данные
                if rows:
                    tabs.append((pymupdf.Rect(t.bbox), rows))
        except Exception:
            pass
        items = [(r.y0, r.x0, "tab", rows) for r, rows in tabs]
        for b in p.get_text("dict")["blocks"]:
            if b["type"] == 1:
                bb = pymupdf.Rect(b["bbox"])
                if bb.width * bb.height < W * H * .85 and bb.width > 40 and bb.height > 30:   # фон-скан листа не берём
                    items.append((bb.y0, bb.x0, "img", bb))
                continue
            for l in b["lines"]:
                bb = pymupdf.Rect(l["bbox"])
                if bb.y1 < H * .08 or bb.y0 > H * .9:        # колонтитулы и номера страниц
                    t = "".join(s["text"] for s in l["spans"])
                    if _is_noise(t) or re.sub(r"\d+", "#", t.strip()) in repeat:
                        continue
                if any((r & bb).get_area() > bb.get_area() * .5 for r, _ in tabs if r.intersects(bb)):
                    continue
                if abs(l["dir"][1]) > .1:                    # повёрнутый текст (штампы, поля)
                    continue
                spans = [s for s in l["spans"] if s["text"].strip()]
                if not spans:
                    items.append((bb.y0, bb.x0, "gap", None))
                    continue
                items.append((bb.y0, bb.x0, "line", {
                    "segs": [(s["text"], _bold(s)) for s in l["spans"]],
                    "x0": spans[0]["bbox"][0], "x1": spans[-1]["bbox"][2], "y0": bb.y0, "y1": bb.y1,
                    "size": max(s["size"] for s in spans),
                    "ital": all(s["flags"] & 2 or "Italic" in s["font"] for s in spans),
                    "col": any(s["color"] & 0xff > 0x80 and s["color"] >> 16 < 0x80 for s in spans)}))
        items.sort(key=lambda i: (round(i[0] / 3), i[1]))
        merged = []      # куски одной строки (текст из нескольких блоков) - в одну строку
        for it in items:
            if it[2] == "line" and merged and merged[-1][2] == "line" and abs(merged[-1][3]["y0"] - it[3]["y0"]) < 2.5 and it[3]["x0"] >= merged[-1][3]["x1"] - 2:
                a, b = merged[-1][3], it[3]
                a["segs"] = a["segs"] + [(" ", False)] + b["segs"]
                a["x1"] = b["x1"]
                continue
            merged.append(it)
        lines = [d for _, _, k, d in merged if k == "line"]
        if lines:
            xs = sorted(d["x0"] for d in lines)
            left = xs[len(xs) // 10]
            right = sorted(d["x1"] for d in lines)[-max(1, len(lines) // 10)]
        else:
            left, right = 0, W
        mid = (left + right) / 2

        out.append(f'<a id="page={pn}"></a>')    # абзац, перешедший со страницы, встанет после якоря - поиск откроет его начало
        for _, _, kind, d in merged:
            if kind == "gap":
                if para:
                    para["gap"] = True
                continue
            if kind == "tab":
                flush()
                rows = d
                if max(len(r) for r in rows) == 1 or len(rows) == 1 and len(rows[0]) <= 2:
                    txt = "".join(f"<p>{esc(c)}</p>" for r in rows for c in r if c)
                    out.append(f'<div class="box">{txt}</div>')
                else:
                    trs = "".join("<tr>" + "".join(f"<td>{esc(c)}</td>" for c in r) + "</tr>" for r in rows)
                    out.append(f'<div class="tbl"><table>{trs}</table></div>')
                continue
            if kind == "img":
                flush()
                nimg += 1
                name = f"{nimg}.png"
                try:
                    p.get_pixmap(clip=d, dpi=150).save(img_dir / name)
                    out.append(f'<figure><img src="{img_url}/{name}" alt="" loading="lazy"></figure>')
                except Exception:
                    pass
                continue
            plain = "".join(t for t, _ in d["segs"]).strip()
            allb = all(b for t, b in d["segs"] if t.strip())
            ind = d["x0"] - left
            centered = ind > 25 and abs((d["x0"] + d["x1"]) / 2 - mid) < 25 and d["x1"] < right - 15
            if d["ital"] and d["col"]:
                k = "ed"
            elif allb and len(plain) < 220 and (centered or plain.isupper() or HEAD_RE.match(plain)) and not (para and para["kind"] == "p" and (right - para["x1"] < 40 or not para["end"]) and ind < 8):
                k = "h"
            elif allb and len(plain) < 220 and para and para["kind"] == "h" and para["page"] == pn and d["y0"] - para["y1"] < para["size"] * 1.6:
                k = "h"      # продолжение заголовка в несколько строк
            elif centered and d["x1"] - d["x0"] < (right - left) * .85:
                k = "c"
            elif ind > (right - left) * .4:
                k = "r"      # гриф «УТВЕРЖДЕНО», адресат - блок справа
            else:
                k = "p"
            same = para is not None and para["page"] == pn
            if para is None or para["kind"] != k:
                new = True
            elif k == "p":     # новый абзац: красная строка, пустая строка, прошлая строка не дошла до края, большой просвет
                # прошлая строка кончилась, хотя первое слово этой строки на ней помещалось - значит, там конец абзаца
                word = len(plain.split()[0]) + 1 if plain.split() else 1
                short = para["end"] and right - para["x1"] > word * para["cw"] + 6
                new = ind > 8 or para["gap"] or short or same and d["y0"] - para["y1"] > para["size"] * 1.1
            else:      # заголовок раздела («I. Общие положения») не приклеиваем к заголовку документа
                new = not same or d["y0"] - para["y1"] > para["size"] * (1.6 if k == "h" else .9) or k == "h" and bool(HEAD_RE.match(plain))
            if new:
                flush()
                para = {"segs": [], "kind": k, "size": d["size"]}
            if k in ("c", "r") and para["segs"]:
                para["segs"].append(("\u2028", False))   # в грифе и строках по центру - разрывы строк как в оригинале
            _add(para["segs"], d["segs"])
            para.update(page=pn, x1=d["x1"], y1=d["y1"], gap=False, end=not plain.endswith("-"),
                        cw=(d["x1"] - d["x0"]) / max(1, len(plain)))    # средняя ширина буквы в этой строке
    flush()
    if not nimg:
        img_dir.rmdir()
    return "\n".join(out)


# ---- текст скана, набранный вручную (documents/_text): «=== стр. N ===», абзац - строка
def sidecar_html(text: str) -> str:
    out = []
    for part in re.split(r"^=== стр\. (\d+) ===\s*$", text, flags=re.M)[1:]:
        if part.isdigit():
            out.append(f'<a id="page={part}"></a>')
            continue
        out += [text_line(x) for x in part.splitlines() if x.strip()]
    return "\n".join(out)


def text_line(x: str) -> str:
    x = re.sub(r"\s+", " ", x).strip()
    if len(x) < 120 and x.isupper():
        return f"<h2>{esc(x)}</h2>"
    return f"<p>{esc(x)}</p>"


def txt_html(path: Path) -> str:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw.decode("cp1251", errors="replace")
    return "\n".join(text_line(x) for x in text.splitlines() if x.strip())


# ---- DOCX: абзацы и таблицы как в документе, оформление - единое
def docx_html(path: Path) -> str:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d, out = docx.Document(path), []
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            p = Paragraph(el, d)
            t = re.sub(r"\s+", " ", p.text).strip()
            if not t:
                continue
            al = str(p.alignment if p.alignment is not None else p.style.paragraph_format.alignment or "")
            runs = [r for r in p.runs if r.text.strip()]
            bold = bool(runs) and all(r.bold or (r.bold is None and p.style.font.bold) for r in runs)
            style = (p.style.name or "").lower()
            if style.startswith(("heading", "заголовок", "title")) or bold and len(t) < 220 and "CENTER" in al:
                lv = heading_level(t)
                out.append(f"<{lv}>{esc(t)}</{lv}>")
            elif "CENTER" in al:
                out.append(f'<p class="c">{esc(t)}</p>' if not t.isupper() else f"<h2>{esc(t)}</h2>")
            elif "RIGHT" in al:
                out.append(f'<p class="r">{esc(t)}</p>')
            elif bold and len(t) < 220 and (HEAD_RE.match(t) or t.isupper()):
                lv = heading_level(t)
                out.append(f"<{lv}>{esc(t)}</{lv}>")
            else:
                out.append(f"<p>{esc(t)}</p>")
        elif el.tag.endswith("}tbl"):
            rows = []
            for r in Table(el, d).rows:
                cells, seen = [], []
                for c in r.cells:
                    if c._tc in seen:          # объединённые ячейки python-docx отдаёт по нескольку раз
                        continue
                    seen.append(c._tc)
                    cells.append("<td>" + "<br>".join(esc(x.text.strip()) for x in c.paragraphs if x.text.strip()) + "</td>")
                rows.append("<tr>" + "".join(cells) + "</tr>")
            out.append('<div class="tbl"><table>' + "".join(rows) + "</table></div>")
    return "\n".join(out)
